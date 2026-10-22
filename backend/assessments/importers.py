"""Question import extraction, parsing and confirmation helpers."""
import csv
import io
import re
import shutil
import subprocess
from pathlib import Path

from django.db import transaction

from . import models

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".jpg", ".jpeg", ".png", ".doc", ".docx", ".xls", ".xlsx", ".csv", ".txt"}


class QuestionImportError(Exception):
    pass


def _spreadsheet_to_text(rows):
    rows = [[str(value).strip() if value is not None else "" for value in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return ""
    headers = [value.lower().replace(" ", "_") for value in rows[0]]
    aliases = {
        "question": {"question", "question_text", "prompt"},
        "type": {"type", "question_type"},
        "marks": {"marks", "mark", "points"},
        "options": {"options", "choices"},
        "answer": {"answer", "correct_answer", "key"},
        "difficulty": {"difficulty", "level"},
        "vocabulary": {"vocabulary", "vocabulary_terms", "vocab", "terms"},
    }
    indexes = {}
    for field, names in aliases.items():
        for index, header in enumerate(headers):
            if header in names:
                indexes[field] = index
                break
    if "question" not in indexes:
        return "\n".join(f"Q: {row[0]}" for row in rows if row and row[0])

    blocks = []
    for row in rows[1:]:
        values = {key: row[index] if index < len(row) else "" for key, index in indexes.items()}
        if not values.get("question"):
            continue
        lines = ["Q: " + values["question"]]
        for key in ("type", "marks", "options", "answer", "difficulty", "vocabulary"):
            if values.get(key):
                lines.append(f"{key.title()}: {values[key]}")
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def extract_file_text(upload):
    """Extract text from supported file types; the original file is retained."""
    if upload is None:
        return "", ""
    filename = Path(upload.name).name
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise QuestionImportError("Supported files: PDF, JPEG/PNG, DOC/DOCX, XLS/XLSX, CSV and TXT.")
    if upload.size > MAX_UPLOAD_BYTES:
        raise QuestionImportError("Maximum upload size is 10 MB.")
    data = upload.read()
    upload.seek(0)
    try:
        if suffix in {".txt", ".csv"}:
            text = data.decode("utf-8-sig", errors="replace")
            return (_spreadsheet_to_text(csv.reader(io.StringIO(text))), "") if suffix == ".csv" else (text, "")
        if suffix == ".pdf":
            from pypdf import PdfReader
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
            return text, "" if text.strip() else "No selectable PDF text found. This may be a scanned PDF; paste its text below."
        if suffix == ".docx":
            from docx import Document
            document = Document(io.BytesIO(data))
            content = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                content.append(_spreadsheet_to_text([[cell.text for cell in row.cells] for row in table.rows]))
            return "\n".join(content), ""
        if suffix == ".doc":
            executable = shutil.which("antiword")
            if executable:
                result = subprocess.run([executable, "-"], input=data, capture_output=True, timeout=15, check=False)
                if result.returncode == 0:
                    return result.stdout.decode("utf-8", errors="replace"), ""
            return "", "Legacy .doc attached, but extraction is unavailable. Save as .docx or paste its text."
        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                return _spreadsheet_to_text(workbook.active.iter_rows(values_only=True)), ""
            finally:
                workbook.close()
        if suffix == ".xls":
            import xlrd
            workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                sheet = workbook.sheet_by_index(0)
                return _spreadsheet_to_text(sheet.row_values(index) for index in range(sheet.nrows)), ""
            finally:
                workbook.release_resources()
        if suffix in {".jpg", ".jpeg", ".png"}:
            try:
                from PIL import Image
                import pytesseract
            except ImportError:
                return "", "Image attached, but OCR is unavailable on this server. Paste the text from the image below."
            try:
                return pytesseract.image_to_string(Image.open(io.BytesIO(data))), ""
            except pytesseract.TesseractNotFoundError:
                return "", "Image attached, but the OCR engine is not installed on this server. Paste the text from the image below."
    except QuestionImportError:
        raise
    except Exception as exc:
        raise QuestionImportError(f"Could not read {filename}: {exc}") from exc
    return "", "Could not extract text from this file. Paste its contents below."


def parse_pasted_questions(text):
    """Parse question blocks (including vocabulary terms) for teacher review."""
    text = (text or "").replace("\r\n", "\n").strip()
    if not text:
        return []
    blocks, current = [], []
    start_pattern = re.compile(r"^\s*(?:Q(?:UESTION)?\s*\d*\s*[:.)-]|\d+[.)])\s*", re.I)
    for line in text.splitlines():
        if start_pattern.match(line) and current and any(part.strip() for part in current):
            blocks.append("\n".join(current).strip())
            current = []
        if not line.strip() and current and any(part.strip() for part in current):
            blocks.append("\n".join(current).strip())
            current = []
        else:
            current.append(line)
    if any(part.strip() for part in current):
        blocks.append("\n".join(current).strip())

    parsed = []
    valid_types = {choice for choice, _ in models.Question.QuestionType.choices}
    valid_difficulties = {choice for choice, _ in models.Question.Difficulty.choices}
    for block in blocks:
        prompt, options, values = [], [], {}
        for raw in block.splitlines():
            line = raw.strip()
            if not line:
                continue
            option = re.match(r"^[A-H][).:-]\s*(.+)$", line, re.I)
            if option:
                options.append(option.group(1).strip())
                continue
            field = re.match(r"^(type|marks|options|answer|correct_answer|difficulty|vocabulary|vocab(?:ulary)?_terms)\s*:\s*(.*)$", line, re.I)
            if field:
                values[field.group(1).lower().replace(" ", "_")] = field.group(2).strip()
                continue
            prompt.append(re.sub(r"^\s*(?:Q(?:UESTION)?\s*\d*\s*[:.)-]|\d+[.)])\s*", "", line, flags=re.I))
        if values.get("options"):
            options.extend(item.strip() for item in re.split(r"\s*\|\s*", values["options"]) if item.strip())
        question_text = " ".join(prompt).strip()
        if not question_text:
            continue
        raw_type = values.get("type", "").upper().replace(" ", "_")
        raw_type = {"MULTIPLE_CHOICE": "MCQ", "TRUEFALSE": "TRUE_FALSE", "SHORT_ANSWER": "SHORT", "LONG_ANSWER": "LONG", "ESSAY": "LONG"}.get(raw_type, raw_type)
        question_type = raw_type if raw_type in valid_types else ("MCQ" if options else "SHORT")
        try:
            marks = max(1, int(values.get("marks", "1") or 1))
        except ValueError:
            marks = 1
        difficulty = values.get("difficulty", "MEDIUM").upper()
        if difficulty not in valid_difficulties:
            difficulty = "MEDIUM"
        vocabulary = values.get("vocabulary", values.get("vocab", values.get("vocabulary_terms", "")))
        parsed.append({
            "question_type": question_type,
            "text": question_text,
            "marks": marks,
            "options": options,
            "correct_answer": values.get("answer", values.get("correct_answer", "")),
            "difficulty": difficulty,
            "vocabulary_terms": [term.strip() for term in re.split(r"[;|]", vocabulary) if term.strip()],
        })
    return parsed


def prepare_question_import(assessment, *, text="", upload=None, actor=None):
    extracted, warning = extract_file_text(upload)
    combined = "\n\n".join(value for value in (extracted, text.strip()) if value)
    parsed = parse_pasted_questions(combined)
    if not parsed and not warning:
        warning = "No question blocks detected. Use the copy/paste template shown on this page."
    imported = models.QuestionImport.objects.create(
        institution_id=assessment.institution_id,
        assessment=assessment,
        source_file=upload,
        original_file_name=Path(upload.name).name if upload else "",
        pasted_text=text,
        extraction_warning=warning,
        parsed_questions=parsed,
        created_by=actor,
    )
    return imported


def refresh_question_import(question_import, *, text="", upload=None):
    if question_import.status == models.QuestionImport.Status.IMPORTED:
        raise QuestionImportError("Imported questions cannot be changed through the preview.")
    extracted, warning = extract_file_text(upload)
    combined = "\n\n".join(value for value in (extracted, text.strip()) if value)
    parsed = parse_pasted_questions(combined)
    if upload:
        question_import.source_file = upload
        question_import.original_file_name = Path(upload.name).name
    question_import.pasted_text = text
    question_import.parsed_questions = parsed
    question_import.extraction_warning = warning or ("No questions detected in the new text." if not parsed else "")
    question_import.save(update_fields=[
        "source_file", "original_file_name", "pasted_text", "parsed_questions",
        "extraction_warning", "updated_at",
    ])
    return question_import


@transaction.atomic
def confirm_question_import(question_import):
    if question_import.status == models.QuestionImport.Status.IMPORTED:
        raise QuestionImportError("This import has already been confirmed.")
    if not question_import.parsed_questions:
        raise QuestionImportError("No questions to import. Paste text or upload a text-readable document.")
    assessment = question_import.assessment
    bank_name = f"{assessment.title} Question Bank"
    bank = models.QuestionBank.all_objects.filter(
        institution_id=assessment.institution_id, subject_id=assessment.subject_id, name=bank_name
    ).first()
    if bank is None:
        bank = models.QuestionBank.all_objects.create(
            institution_id=assessment.institution_id, name=bank_name,
            subject_id=assessment.subject_id, description=f"Imported for {assessment.title}",
        )
    for item in question_import.parsed_questions:
        question = models.Question.all_objects.create(
            institution_id=assessment.institution_id, bank=bank,
            question_type=item["question_type"], text=item["text"], marks=item["marks"],
            options=item.get("options", []), correct_answer=item.get("correct_answer", ""),
            difficulty=item.get("difficulty", "MEDIUM"),
            vocabulary_terms=item.get("vocabulary_terms", []),
        )
        models.AssessmentQuestion.all_objects.create(
            institution_id=assessment.institution_id, assessment=assessment,
            question=question, marks_override=question.marks,
        )
    question_import.status = models.QuestionImport.Status.IMPORTED
    question_import.imported_count = len(question_import.parsed_questions)
    question_import.save(update_fields=["status", "imported_count", "updated_at"])
    return question_import
