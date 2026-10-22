"""Assessments API tests."""
import pytest
from io import BytesIO
from tempfile import TemporaryDirectory
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from PIL import Image

from academics.models import AcademicYear, GradeLevel, SchoolClass, Section, Subject
from people import services as people_services

pytestmark = pytest.mark.django_db


@pytest.fixture
def section(institution):
    grade = GradeLevel.objects.create(institution=institution, name="Grade 10")
    school_class = SchoolClass.objects.create(institution=institution, grade_level=grade, name="Grade 10")
    return Section.objects.create(institution=institution, school_class=school_class, name="A")


@pytest.fixture
def academic_year(institution):
    return AcademicYear.objects.create(
        institution=institution, name="2026-2027",
        start_date="2026-06-01", end_date="2027-03-31", is_active=True,
    )


@pytest.fixture
def subject(institution):
    return Subject.objects.create(institution=institution, code="MATH", name="Mathematics")


@pytest.fixture
def students(institution, section):
    return [
        people_services.create_student(institution.id, f"s{i}@test.school", roll_number=f"R{i}", admission_number=f"ADM{i}", section=section)
        for i in range(1, 4)
    ]


@pytest.fixture
def assessment(api_client, subject, section, academic_year):
    response = api_client.post(
        "/api/assessments/",
        {
            "title": "Midterm",
            "assessment_type": "TEST",
            "subject": subject.id,
            "section": section.id,
            "academic_year": academic_year.id,
            "date": "2026-10-26",
            "total_marks": 100,
            "pass_marks": 40,
        },
        format="json",
    )
    assert response.status_code == 201, response.data
    return response.data["data"]


def test_create_assessment(assessment):
    assert assessment["title"] == "Midterm"
    assert assessment["total_marks"] == 100


def test_assign_students_to_assessment(api_client, assessment, students):
    response = api_client.post(
        f"/api/assessments/{assessment['id']}/assign/",
        {"students": [students[0].id, students[2].id]},
        format="json",
    )
    assert response.status_code == 200, response.data
    assert response.data["data"] == {"created": 2, "already_assigned": 0, "total": 2}

    listing = api_client.get(f"/api/student-assessments/?assessment={assessment['id']}")
    assert listing.status_code == 200
    assert listing.data["data"]["total"] == 2
    assert all(row["status"] == "PENDING" for row in listing.data["data"]["items"])

    repeated = api_client.post(
        f"/api/assessments/{assessment['id']}/assign/",
        {"students": [students[0].id]},
        format="json",
    )
    assert repeated.data["data"]["already_assigned"] == 1


def test_assigned_assessment_is_visible_to_that_student_only(api_client, assessment, students):
    from rest_framework.test import APIClient

    api_client.post(
        f"/api/assessments/{assessment['id']}/assign/",
        {"students": [students[0].id]},
        format="json",
    )
    student_client = APIClient()
    student_client.force_authenticate(user=students[0].user)
    mine = student_client.get("/api/student-assessments/mine/")
    assert mine.status_code == 200, mine.data
    assert len(mine.data["data"]) == 1
    assert mine.data["data"][0]["assessment_detail"]["title"] == assessment["title"]

    another = people_services.create_student(
        students[0].institution_id, "unassigned@test.school", roll_number="U1", admission_number="AU1", section=students[0].section
    )
    another_client = APIClient()
    another_client.force_authenticate(user=another.user)
    assert another_client.get("/api/student-assessments/mine/").data["data"] == []


def test_student_can_view_safe_questions_and_submit_answers(api_client, assessment, students, subject):
    from rest_framework.test import APIClient

    bank = api_client.post(
        "/api/question-banks/", {"name": "Student quiz bank", "subject": subject.id}, format="json"
    )
    question = api_client.post(
        f"/api/question-banks/{bank.data['data']['id']}/questions/",
        {
            "question_type": "MCQ", "text": "2 + 2 = ?", "marks": 1,
            "options": ["3", "4"], "correct_answer": "4", "difficulty": "EASY",
        },
        format="json",
    )
    qid = question.data["data"]["id"]
    api_client.post(
        "/api/assessment-questions/",
        {"assessment": assessment["id"], "question": qid, "marks_override": 1},
        format="json",
    )
    api_client.post(
        f"/api/assessments/{assessment['id']}/assign/",
        {"students": [students[0].id]},
        format="json",
    )

    student_client = APIClient()
    student_client.force_authenticate(user=students[0].user)
    mine = student_client.get("/api/student-assessments/mine/")
    assert mine.status_code == 200, mine.data
    row = mine.data["data"][0]
    assert row["questions"][0]["text"] == "2 + 2 = ?"
    assert "correct_answer" not in row["questions"][0]

    submitted = student_client.post(
        f"/api/student-assessments/{row['id']}/submit/",
        {"answers": {str(qid): "4"}},
        format="json",
    )
    assert submitted.status_code == 200, submitted.data
    assert submitted.data["data"]["status"] == "GRADED"
    assert submitted.data["data"]["obtained_marks"] == "1.00"
    assert submitted.data["data"]["answers"][str(qid)] == "4"
    second_submit = student_client.post(
        f"/api/student-assessments/{row['id']}/submit/",
        {"answers": {str(qid): "3"}},
        format="json",
    )
    assert second_submit.status_code == 400


def test_objective_only_submission_is_auto_graded(api_client, assessment, students, subject):
    from rest_framework.test import APIClient

    bank = api_client.post(
        "/api/question-banks/", {"name": "Objective bank", "subject": subject.id}, format="json"
    )
    question = api_client.post(
        f"/api/question-banks/{bank.data['data']['id']}/questions/",
        {"question_type": "TRUE_FALSE", "text": "Water is wet.", "marks": 2, "options": ["True", "False"], "correct_answer": "True", "difficulty": "EASY"},
        format="json",
    )
    qid = question.data["data"]["id"]
    api_client.post("/api/assessment-questions/", {"assessment": assessment["id"], "question": qid, "marks_override": 2}, format="json")
    api_client.post(f"/api/assessments/{assessment['id']}/assign/", {"students": [students[0].id]}, format="json")

    client = APIClient()
    client.force_authenticate(user=students[0].user)
    row = client.get("/api/student-assessments/mine/").data["data"][0]
    submitted = client.post(
        f"/api/student-assessments/{row['id']}/submit/",
        {"answers": {str(qid): "True"}}, format="json",
    )
    assert submitted.status_code == 200, submitted.data
    assert submitted.data["data"]["status"] == "GRADED"
    assert submitted.data["data"]["obtained_marks"] == "2.00"


def test_assign_rejects_student_from_other_section(api_client, assessment, institution, section):
    from academics.models import Section

    other_section = Section.objects.create(
        institution=institution, school_class=section.school_class, name="B"
    )
    student = people_services.create_student(
        institution.id, "assign-other-section@test.school", roll_number="F1",
        admission_number="AF1", section=other_section,
    )
    response = api_client.post(
        f"/api/assessments/{assessment['id']}/assign/",
        {"students": [student.id]},
        format="json",
    )
    assert response.status_code == 400
    marks = api_client.post(
        f"/api/assessments/{assessment['id']}/marks/",
        {"records": [{"student": student.id, "obtained_marks": "5", "status": "GRADED"}]},
        format="json",
    )
    assert marks.status_code == 400


def test_bulk_marks_and_grades(api_client, assessment, students):
    # grade scale
    api_client.post("/api/grade-scales/", {"grade": "A", "min_mark": 80, "max_mark": 100, "remark": "Excellent"}, format="json")
    api_client.post("/api/grade-scales/", {"grade": "B", "min_mark": 60, "max_mark": 79, "remark": "Good"}, format="json")
    api_client.post("/api/grade-scales/", {"grade": "F", "min_mark": 0, "max_mark": 39, "remark": "Fail"}, format="json")

    response = api_client.post(
        f"/api/assessments/{assessment['id']}/marks/",
        {
            "records": [
                {"student": students[0].id, "obtained_marks": "85", "status": "GRADED"},
                {"student": students[1].id, "obtained_marks": "55", "status": "GRADED"},
                {"student": students[2].id, "obtained_marks": "30", "status": "GRADED"},
            ]
        },
        format="json",
    )
    assert response.status_code == 200, response.data
    assert response.data["data"]["created"] == 3


def test_results_ranking_and_summary(api_client, assessment, students):
    api_client.post("/api/grade-scales/", {"grade": "A", "min_mark": 80, "max_mark": 100}, format="json")
    api_client.post("/api/grade-scales/", {"grade": "F", "min_mark": 0, "max_mark": 39}, format="json")
    api_client.post(
        f"/api/assessments/{assessment['id']}/marks/",
        {
            "records": [
                {"student": students[0].id, "obtained_marks": "90", "status": "GRADED"},
                {"student": students[1].id, "obtained_marks": "50", "status": "GRADED"},
            ]
        },
        format="json",
    )
    response = api_client.get(f"/api/assessments/{assessment['id']}/results/")
    assert response.status_code == 200, response.data
    data = response.data["data"]
    assert data["summary"]["total"] == 2
    assert data["summary"]["passed"] == 2  # pass_marks 40; both 90 and 50 pass
    # grades assigned via scale
    top = data["results"][0]
    assert top["grade"] == "A"
    assert top["rank"] == 1


def test_question_bank_and_questions(api_client, subject):
    bank = api_client.post("/api/question-banks/", {"name": "Bank 1", "subject": subject.id}, format="json")
    assert bank.status_code == 201, bank.data
    bank_id = bank.data["data"]["id"]

    q = api_client.post(
        f"/api/question-banks/{bank_id}/questions/",
        {"question_type": "MCQ", "text": "2+2?", "marks": 1, "options": ["3", "4"], "correct_answer": "4", "difficulty": "EASY"},
        format="json",
    )
    assert q.status_code == 201, q.data

    listed = api_client.get(f"/api/question-banks/{bank_id}/questions/")
    assert len(listed.data["data"]) == 1

    banks = api_client.get("/api/question-banks/")
    assert banks.data["data"]["items"][0]["question_count"] == 1


def test_tenant_isolation(api_client, academic_year, other_institution):
    from academics.models import Subject

    foreign_subject = Subject.objects.create(institution=other_institution, code="PHY", name="Physics")
    response = api_client.post(
        "/api/assessments/",
        {
            "title": "Foreign",
            "assessment_type": "TEST",
            "subject": foreign_subject.id,
            "academic_year": academic_year.id,
            "date": "2026-10-26",
            "total_marks": 100,
        },
        format="json",
    )
    assert response.status_code == 400


def test_copy_paste_question_preview_and_confirm(api_client, assessment):
    response = api_client.post(
        f"/api/assessments/{assessment['id']}/question-imports/preview/",
        {"text": "Q: Define evaporation?\nType: SHORT\nMarks: 2\nVocabulary: evaporation: liquid to gas; vapour: gas form"},
        format="json",
    )
    assert response.status_code == 201, response.data
    preview = response.data["data"]
    assert preview["status"] == "PREVIEW"
    assert preview["parsed_questions"][0]["text"] == "Define evaporation?"
    assert preview["parsed_questions"][0]["vocabulary_terms"] == [
        "evaporation: liquid to gas", "vapour: gas form"
    ]

    confirmed = api_client.post(f"/api/question-imports/{preview['id']}/confirm/")
    assert confirmed.status_code == 200, confirmed.data
    assert confirmed.data["data"]["imported_count"] == 1

    from assessments.models import AssessmentQuestion

    link = AssessmentQuestion.objects.get(assessment_id=assessment["id"])
    assert link.question.vocabulary_terms[0].startswith("evaporation")

    duplicate_confirm = api_client.post(f"/api/question-imports/{preview['id']}/confirm/")
    assert duplicate_confirm.status_code == 400


def test_csv_question_file_preview(api_client, assessment):
    upload = SimpleUploadedFile(
        "questions.csv",
        b"Question,Type,Marks,Options,Answer,Difficulty,Vocabulary\n"
        b"Which is a prime number?,MCQ,1,4 | 5 | 6,5,EASY,prime: divisible by itself and one\n",
        content_type="text/csv",
    )
    with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
        response = api_client.post(
            f"/api/assessments/{assessment['id']}/question-imports/preview/",
            {"file": upload},
            format="multipart",
        )
        assert response.status_code == 201, response.data
        imported = response.data["data"]
        assert imported["original_file_name"] == "questions.csv"
        assert imported["parsed_questions"][0]["options"] == ["4", "5", "6"]


def test_jpeg_attachment_is_accepted_with_manual_transcription_fallback(api_client, assessment):
    image_bytes = BytesIO()
    Image.new("RGB", (16, 16), color="white").save(image_bytes, format="JPEG")
    upload = SimpleUploadedFile("paper.jpg", image_bytes.getvalue(), content_type="image/jpeg")
    with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
        response = api_client.post(
            f"/api/assessments/{assessment['id']}/question-imports/preview/",
            {"file": upload},
            format="multipart",
        )
        assert response.status_code == 201, response.data
        assert response.data["data"]["original_file_name"] == "paper.jpg"
        assert response.data["data"]["extraction_warning"]
        assert response.data["data"]["parsed_questions"] == []
        import_id = response.data["data"]["id"]
        refreshed = api_client.post(
            f"/api/question-imports/{import_id}/refresh/",
            {"text": "Q: What is a habitat?\nVocabulary: habitat: an organism's home"},
            format="json",
        )
        assert refreshed.status_code == 200, refreshed.data
        assert refreshed.data["data"]["original_file_name"] == "paper.jpg"
        assert refreshed.data["data"]["parsed_questions"][0]["vocabulary_terms"] == [
            "habitat: an organism's home"
        ]


def test_excel_question_upload(api_client, assessment):
    from openpyxl import Workbook

    content = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Question", "Type", "Marks", "Options", "Answer", "Vocabulary"])
    sheet.append(["2 + 2 equals?", "MCQ", 1, "3 | 4", "4", "sum: result of addition"])
    workbook.save(content)
    upload = SimpleUploadedFile("questions.xlsx", content.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
        response = api_client.post(
            f"/api/assessments/{assessment['id']}/question-imports/preview/",
            {"file": upload}, format="multipart",
        )
        assert response.status_code == 201, response.data
        question = response.data["data"]["parsed_questions"][0]
        assert question["text"] == "2 + 2 equals?"
        assert question["vocabulary_terms"] == ["sum: result of addition"]


def test_docx_question_upload(api_client, assessment):
    from docx import Document

    content = BytesIO()
    document = Document()
    document.add_paragraph("Q: Define gravity")
    document.add_paragraph("Type: SHORT")
    document.add_paragraph("Vocabulary: gravity: force pulling objects together")
    document.save(content)
    upload = SimpleUploadedFile("questions.docx", content.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
        response = api_client.post(
            f"/api/assessments/{assessment['id']}/question-imports/preview/",
            {"file": upload}, format="multipart",
        )
        assert response.status_code == 201, response.data
        question = response.data["data"]["parsed_questions"][0]
        assert question["text"] == "Define gravity"
        assert question["vocabulary_terms"] == ["gravity: force pulling objects together"]


def test_question_import_works_without_assessment_subject(api_client, academic_year):
    response = api_client.post(
        "/api/assessments/",
        {"title": "No subject", "academic_year": academic_year.id, "date": "2026-10-26"},
        format="json",
    )
    assert response.status_code == 201, response.data
    assessment_id = response.data["data"]["id"]
    preview = api_client.post(
        f"/api/assessments/{assessment_id}/question-imports/preview/",
        {"text": "Q: What is water?"},
        format="json",
    )
    assert preview.status_code == 201, preview.data
    confirmed = api_client.post(f"/api/question-imports/{preview.data['data']['id']}/confirm/")
    assert confirmed.status_code == 200, confirmed.data


def test_invalid_import_extension_is_rejected(api_client, assessment):
    upload = SimpleUploadedFile("malware.exe", b"not executable", content_type="application/octet-stream")
    response = api_client.post(
        f"/api/assessments/{assessment['id']}/question-imports/preview/",
        {"file": upload},
        format="multipart",
    )
    assert response.status_code == 400
